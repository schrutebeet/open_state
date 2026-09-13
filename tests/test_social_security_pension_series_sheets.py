from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook

from civic_metrics.catalog import load_catalog
from civic_metrics.connectors.social_security_pensions import SocialSecurityPensionsConnector
from civic_metrics.domain import DatasetPayload

ROOT = Path(__file__).parents[1]


def _payload(workbook: Workbook, filename: str = "CA202601.xlsx") -> DatasetPayload:
    stream = BytesIO()
    workbook.save(stream)
    body = stream.getvalue()
    return DatasetPayload(
        dataset_code="social_security_pension_series",
        source_code="social_security",
        fetched_at=datetime.now(UTC),
        source_url=f"https://example.test/{filename}",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        body=body,
        sha256="test",
        metadata={},
    )


def _catalog_items():
    catalog = load_catalog(ROOT / "config")
    dataset = catalog.dataset_by_code["social_security_pension_series"]
    indicators = [item for item in catalog.indicators if item.dataset == dataset.code]
    return dataset, indicators


def _national_sheet(title: str, pension_count: int) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(["Territorio", "TOTAL PENSIONES", None, "JUBILACIÓN"])
    sheet.append([None, "Número", "P.media", "Número", "P.media"])
    sheet.append(["TOTAL", pension_count, 1250, 600, 1400])
    return workbook


def test_ca2_is_used_when_primary_sheet_is_missing_and_precedes_legacy_sheet() -> None:
    workbook = _national_sheet("CA2", 1000)
    legacy = workbook.create_sheet("TOTALSISTEMA.")
    legacy.append(["Territorio", "TOTAL PENSIONES", None, "JUBILACIÓN"])
    legacy.append([None, "Número", "P.media", "Número", "P.media"])
    legacy.append(["TOTAL", 9999, 999, 999, 999])
    dataset, indicators = _catalog_items()

    candidates = SocialSecurityPensionsConnector().extract(dataset, _payload(workbook), indicators)

    values = {candidate.indicator_code: candidate.value for candidate in candidates}
    assert values == {
        "pension_count": 1000,
        "average_pension": 1250,
        "average_retirement_pension": 1400,
    }


@pytest.mark.parametrize("sheet_name", ["TOTALSISTEMA.", "TOTALSISTEMA", "Total Sistema"])
def test_legacy_sheet_aggregates_all_regions_with_weighted_averages(sheet_name: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(["1 de enero de 2016"])
    sheet.append(
        [
            "Territorio",
            "INCAPACIDAD PERMANENTE",
            None,
            "JUBILACIÓN",
            None,
            "VIUDEDAD",
            None,
            "ORFANDAD",
            None,
            "FAVOR FAMILIARES",
            None,
            "TOTAL PENSIONES",
            None,
        ]
    )
    sheet.append(
        [
            None,
            "Número",
            "P.media",
            "Número",
            "P.media",
            "Número",
            "P.media",
            "Número",
            "P.media",
            "Número",
            "P.media",
            "Número",
            "P.media",
        ]
    )

    areas = [
        "ANDALUCÍA",
        "ARAGÓN",
        "ASTURIAS (PRINCIPADO DE)",
        "BALEARS (ILLES)",
        "CANARIAS",
        "CANTABRIA",
        "CASTILLA - LA MANCHA",
        "CASTILLA Y LEÓN",
        "CATALUÑA",
        "COMUNITAT VALENCIANA",
        "EXTREMADURA",
        "GALICIA",
        "MADRID (COM. DE)",
        "MURCIA (REGIÓN DE)",
        "NAVARRA (COM. FORAL DE)",
        "PAÍS VASCO",
        "RIOJA (LA)",
        "Ceuta",
        "Melilla",
    ]
    for index, area in enumerate(areas):
        row = [area, 1, 100, 40, 1200 + index, 1, 100, 1, 100, 1, 100]
        row.extend([100, 2000 + index])
        sheet.append(row)
    dataset, indicators = _catalog_items()

    candidates = SocialSecurityPensionsConnector().extract(
        dataset, _payload(workbook, "CA2(012016).xlsx"), indicators
    )

    values = {candidate.indicator_code: candidate.value for candidate in candidates}
    assert values["pension_count"] == 1900
    assert values["average_pension"] == 2009
    assert values["average_retirement_pension"] == 1209
    assert {candidate.period.label for candidate in candidates} == {"2016-01"}
    assert all("19 areas" in candidate.source_series for candidate in candidates)


def test_legacy_ca2_filename_provides_month_when_workbook_has_no_date() -> None:
    payload = _payload(Workbook(), "CA2%28012016%29.xlsx")

    period = SocialSecurityPensionsConnector._period_from_payload(payload)

    assert period.label == "2016-01"
