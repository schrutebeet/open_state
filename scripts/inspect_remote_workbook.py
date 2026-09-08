import argparse
from io import BytesIO
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from civic_metrics.http import HttpClient

parser = argparse.ArgumentParser()
parser.add_argument('url')
parser.add_argument('--sheet')
parser.add_argument('--rows', type=int, default=12)
args = parser.parse_args()
http = HttpClient(30)
try:
    response = http.get(args.url)
    book = openpyxl.load_workbook(BytesIO(response.body), read_only=True, data_only=True)
    print('SHEETS', book.sheetnames)
    for sheet in book:
        if args.sheet and args.sheet != sheet.title:
            continue
        print('SHEET', sheet.title, sheet.max_row, sheet.max_column)
        for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
            if i > args.rows:
                break
            print(i, list(row)[:18])
    book.close()
finally:
    http.close()
