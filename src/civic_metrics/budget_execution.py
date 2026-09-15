"""Update and persist IGAE's monthly State budget execution archive."""

from __future__ import annotations

import calendar
import json
import logging
import os
import re
import sqlite3
import tempfile
import unicodedata
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote, urljoin, urlparse

import httpx
import openpyxl
import xlrd
from bs4 import BeautifulSoup

LOGGER = logging.getLogger(__name__)

ARCHIVE_URL = (
    "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/"
    "ContabilidadPublica/CPE/EjecucionPresupuestaria/Paginas/imoperacionesejecucion.aspx"
)
MONTHS = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10,
    "noviembre": 11, "diciembre": 12,
}
FIELDS = [
    "initial_credits_thousand_eur", "credit_modifications_thousand_eur",
    "final_credits_thousand_eur", "committed_expenses_thousand_eur",
    "recognized_obligations_thousand_eur", "payments_thousand_eur",
    "pending_payments_thousand_eur", "credit_remnants_thousand_eur",
]
FLOW_FIELDS = [
    "credit_modifications_thousand_eur", "committed_expenses_thousand_eur",
    "recognized_obligations_thousand_eur", "payments_thousand_eur",
]
LABEL_RE = re.compile(r"^\s*(\d+)\.\s*(.+?)\s*$")


def _simplify(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text).strip().lower()


def _amount(value: object) -> int | float:
    if value is None or str(value).strip() in {"", "-", "–", "—"}:
        return 0
    if isinstance(value, (int, float)):
        number = Decimal(str(value))
    else:
        text = str(value).strip().replace("\xa0", "").replace(" ", "")
        negative = text.startswith("(") and text.endswith(")")
        text = text.strip("()").replace("−", "-").replace("–", "-").replace("—", "-")
        if "," in text:
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(".", "")
        try:
            number = Decimal(text)
        except InvalidOperation as exc:
            raise ValueError(f"Invalid IGAE amount {value!r}") from exc
        if negative:
            number = -number
    return int(number) if number == number.to_integral_value() else float(number)


def _mapped(values: list[object]) -> dict[str, int | float]:
    if len(values) != len(FIELDS):
        raise ValueError(f"Expected {len(FIELDS)} financial columns, got {len(values)}")
    return dict(zip(FIELDS, (_amount(value) for value in values), strict=True))


def _period_from_text(text: str) -> str | None:
    normal = _simplify(unquote(text))
    year_match = re.search(r"\b(20\d{2})\b", normal)
    if not year_match:
        return None
    year = year_match.group(1)
    for name, month in sorted(MONTHS.items(), key=lambda item: -len(item[0])):
        if re.search(rf"(?<![a-z]){name}(?![a-z])", normal):
            return f"{year}-{month:02d}"
    compact = re.search(r"\b(20\d{2})(0[1-9]|1[0-2])\b", normal)
    return f"{compact.group(1)}-{compact.group(2)}" if compact else None


def _candidate_rank(item: dict) -> tuple[int, int]:
    text = _simplify(f"{item['text']} {item['url']}")
    rank = {"xlsx": 30, "xls": 20, "pdf": 10}.get(item["format"], 0)
    if "cuadros" in text or "excel" in text:
        rank += 10
    return rank, -len(item["url"])


def _discover_current_sources(client: httpx.Client, year: int) -> dict[str, dict]:
    response = client.get(ARCHIVE_URL)
    response.raise_for_status()
    soup = BeautifulSoup(response.content, "html.parser")
    candidates: dict[str, list[dict]] = {}
    for anchor in soup.find_all("a", href=True):
        url = urljoin(str(response.url), str(anchor["href"]))
        path = unquote(urlparse(url).path).lower()
        extension = Path(path).suffix.lstrip(".")
        if extension not in {"xls", "xlsx", "pdf"}:
            continue
        text = " ".join(anchor.stripped_strings)
        period = _period_from_text(f"{path} {text}")
        if not period or int(period[:4]) != year:
            continue
        item = {
            "url": url, "archive_url": str(response.url), "format": extension, "text": text,
        }
        candidates.setdefault(period, []).append(item)
    return {period: max(items, key=_candidate_rank) for period, items in candidates.items()}


def _rows_from_workbook(body: bytes, extension: str) -> tuple[str, list[list[object]]]:
    if extension == "xls" and body[:2] != b"PK":
        workbook = xlrd.open_workbook(file_contents=body)
        names = workbook.sheet_names()
        name = next((n for n in names if _simplify(n) in {"gtos 008", "egb02"}), None)
        if not name:
            raise LookupError(f"No GTOS 008/EGB02 sheet; sheets={names}")
        sheet = workbook.sheet_by_name(name)
        return name, [sheet.row_values(index) for index in range(sheet.nrows)]
    workbook = openpyxl.load_workbook(BytesIO(body), data_only=True, read_only=True)
    try:
        name = next((n for n in workbook.sheetnames if _simplify(n) in {"gtos 008", "egb02"}), None)
        if not name:
            raise LookupError(f"No GTOS 008/EGB02 sheet; sheets={workbook.sheetnames}")
        return name, [list(row) for row in workbook[name].iter_rows(values_only=True)]
    finally:
        workbook.close()


def _parse_source(body: bytes, extension: str, period: str, source: dict, metadata_by_code: dict) -> dict:
    if extension == "pdf":
        raise ValueError("PDF is not accepted for the live updater; an Excel 'Cuadros' source is required")
    sheet_name, rows = _rows_from_workbook(body, extension)
    categories: list[dict] = []
    totals = None
    parent = None
    for row_number, row in enumerate(rows, start=1):
        if not row or row[0] in (None, ""):
            continue
        label = str(row[0]).replace("\n", " ").strip()
        if _simplify(label) == "totales":
            if len(row) >= 9:
                totals = _mapped(list(row[1:9]))
            continue
        match = LABEL_RE.match(label)
        if not match or len(row) < 9:
            continue
        local_code, name = match.groups()
        name = name.strip()
        is_parent = name.isupper() and any(char.isalpha() for char in name)
        if is_parent or parent is None:
            code, level, parent_code = local_code, "area", None
            parent = {"code": code, "name_es": name}
        else:
            code, level, parent_code = f"{parent['code']}.{local_code}", "policy", parent["code"]
        known = metadata_by_code.get((level, code), {})
        categories.append({
            "code": code, "source_code": local_code, "name_es": name,
            "name_full_es": known.get("name_full_es", name),
            "description_es": known.get("description_es", f"Partida presupuestaria de {name.lower()}."),
            "name": known.get("name", name), "name_full": known.get("name_full", name),
            "description": known.get("description", f"Budget category for {name.lower()}."),
            "level": level, "parent_code": parent_code, "amounts": _mapped(list(row[1:9])),
            "source_row": row_number,
        })
    if not categories or totals is None:
        raise ValueError(f"{sheet_name} contains no parseable categories and totals")
    return {
        "period": period, "period_semantics": "Budget status and cumulative execution through the end of this month",
        "classification": "areas_and_policies", "table_or_sheet": sheet_name,
        "source_unit": "thousand_eur", "field_order_in_source": FIELDS,
        "categories": categories, "totals": totals,
    }


def _validate(record: dict) -> dict:
    categories = record["categories"]
    totals = record["totals"]
    checks: dict[str, bool] = {}
    for subset_name, subset in (
        ("aggregate_rows_sum_to_total", [c for c in categories if c["parent_code"] is None]),
        ("detail_rows_sum_to_total", [c for c in categories if c["parent_code"] is not None]),
    ):
        checks[subset_name] = {
            field: sum(Decimal(str(c["amounts"][field])) for c in subset) == Decimal(str(totals[field]))
            for field in FIELDS
        }
    checks["recognized_obligations_equal_payments_plus_pending"] = (
        totals["recognized_obligations_thousand_eur"]
        == totals["payments_thousand_eur"] + totals["pending_payments_thousand_eur"]
    )
    checks["final_credits_equal_obligations_plus_credit_remnants"] = (
        totals["final_credits_thousand_eur"]
        == totals["recognized_obligations_thousand_eur"] + totals["credit_remnants_thousand_eur"]
    )
    checks["all_financial_columns_reconcile"] = all(
        value is True for group in checks.values() if isinstance(group, dict) for value in group.values()
    ) and all(value is True for value in checks.values() if isinstance(value, bool))
    return checks


def _previous_period(period: str) -> str | None:
    year, month = map(int, period.split("-"))
    return None if month == 1 else f"{year}-{month - 1:02d}"


def _derive_monthly(periods: dict[str, dict]) -> None:
    for period, record in sorted(periods.items()):
        previous = None if period.endswith("-01") else periods.get(_previous_period(period))
        is_january = period.endswith("-01")
        if record.get("data_completeness") != "complete" or (previous is None and not is_january):
            status, reason = "unavailable", "Missing previous month or partial source"
        else:
            status, reason = "derived", None
        source_urls = [record["reference_url"]]
        if previous:
            source_urls.insert(0, previous["reference_url"])
        current_by_key = {(c["level"], c["code"]): c for c in record["categories"]}
        previous_by_key = {(c["level"], c["code"]): c for c in (previous or {}).get("categories", [])}
        for key, category in current_by_key.items():
            prior = previous_by_key.get(key)
            flows = {field: None for field in FLOW_FIELDS}
            balances = {"pending_payments_thousand_eur": None, "credit_remnants_thousand_eur": None}
            if status == "derived":
                for field in FLOW_FIELDS:
                    current = category["amounts"][field]
                    flows[field] = current if prior is None else current - prior["amounts"][field]
                for field in balances:
                    balances[field] = category["amounts"][field] - prior["amounts"][field] if prior else None
            category["monthly_flows"] = {
                "status": status, "flows": flows,
                "budget_changes": {
                    "net_change_in_final_credits_thousand_eur": (
                        category["amounts"]["final_credits_thousand_eur"]
                        if status == "derived" and prior is None
                        else category["amounts"]["final_credits_thousand_eur"] - prior["amounts"]["final_credits_thousand_eur"]
                        if status == "derived" else None
                    )
                },
                "balance_changes": balances,
            }
        totals = record["totals"]
        prior_totals = (previous or {}).get("totals", {})
        record["monthly_flows"] = {
            "status": status,
            "method": "January: use January year-to-date; later months: current YTD minus prior month YTD within the year.",
            "derived_from_periods": [p for p in ([_previous_period(period), period]) if p and p in periods],
            "source_urls": source_urls, "reason": reason,
            "flows": {
                field: totals[field] if status == "derived" and previous is None else
                totals[field] - prior_totals[field] if status == "derived" else None
                for field in FLOW_FIELDS
            },
            "budget_changes": {
                "net_change_in_final_credits_thousand_eur": totals["final_credits_thousand_eur"] if status == "derived" and previous is None else
                totals["final_credits_thousand_eur"] - prior_totals["final_credits_thousand_eur"] if status == "derived" else None,
            },
            "balance_changes": {
                field: totals[field] if status == "derived" and previous is None else
                totals[field] - prior_totals[field] if status == "derived" else None
                for field in ("pending_payments_thousand_eur", "credit_remnants_thousand_eur")
            },
            "category_coverage": {
                "matched_current_category_count": len(set(current_by_key) & set(previous_by_key)) if previous else len(current_by_key),
                "unmatched_current_category_codes": sorted(f"{level}:{code}" for level, code in set(current_by_key) - set(previous_by_key)) if previous else [],
                "previous_only_category_codes": sorted(f"{level}:{code}" for level, code in set(previous_by_key) - set(current_by_key)) if previous else [],
                "complete": not previous or set(current_by_key) == set(previous_by_key),
            },
        }


def _create_db_tables(db: sqlite3.Connection) -> None:
    db.executescript("""
    CREATE TABLE IF NOT EXISTS state_budget_metadata (key TEXT PRIMARY KEY, value_json TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS state_budget_periods (
      period TEXT PRIMARY KEY, period_start TEXT NOT NULL, period_end TEXT NOT NULL,
      period_semantics TEXT NOT NULL, classification TEXT NOT NULL, source_unit TEXT NOT NULL,
      source_format TEXT NOT NULL, source_file TEXT NOT NULL, reference_url TEXT NOT NULL,
      archive_page_url TEXT, field_order_json TEXT NOT NULL, source_record_json TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS state_budget (
      id INTEGER PRIMARY KEY AUTOINCREMENT, period TEXT NOT NULL, level TEXT NOT NULL, code TEXT NOT NULL,
      source_code TEXT NOT NULL, parent_code TEXT, name_es TEXT NOT NULL, name_full_es TEXT,
      description_es TEXT, name TEXT, name_full TEXT, description TEXT,
      initial_credits_thousand_eur NUMERIC, credit_modifications_thousand_eur NUMERIC,
      final_credits_thousand_eur NUMERIC, classification TEXT NOT NULL, source_url TEXT NOT NULL,
      source_file TEXT NOT NULL, source_format TEXT NOT NULL, source_row_json TEXT NOT NULL,
      UNIQUE(period, level, code), FOREIGN KEY(period) REFERENCES state_budget_periods(period)
    );
    CREATE TABLE IF NOT EXISTS state_expenditure (
      id INTEGER PRIMARY KEY AUTOINCREMENT, budget_id INTEGER, period TEXT NOT NULL, level TEXT NOT NULL,
      code TEXT NOT NULL, committed_expenses_thousand_eur NUMERIC,
      recognized_obligations_thousand_eur NUMERIC, payments_thousand_eur NUMERIC,
      pending_payments_thousand_eur NUMERIC, credit_remnants_thousand_eur NUMERIC,
      monthly_committed_expenses_thousand_eur NUMERIC, monthly_recognized_obligations_thousand_eur NUMERIC,
      monthly_payments_thousand_eur NUMERIC, monthly_credit_modifications_thousand_eur NUMERIC,
      monthly_pending_payments_change_thousand_eur NUMERIC, monthly_credit_remnants_change_thousand_eur NUMERIC,
      monthly_status TEXT NOT NULL, monthly_coverage_json TEXT NOT NULL, source_url TEXT NOT NULL,
      UNIQUE(period, level, code), FOREIGN KEY(budget_id) REFERENCES state_budget(id),
      FOREIGN KEY(period) REFERENCES state_budget_periods(period)
    );
    CREATE TABLE IF NOT EXISTS state_budget_audit (
      id INTEGER PRIMARY KEY CHECK(id=1), summary_json TEXT NOT NULL, imported_at TEXT NOT NULL
    );
    """)


def _j(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _upsert_db(db_path: Path, data: dict, audit: dict) -> None:
    with sqlite3.connect(db_path) as db:
        db.execute("PRAGMA foreign_keys=ON")
        _create_db_tables(db)
        metadata = {key: data.get(key) for key in (
            "dataset", "title_es", "institution", "scope", "requested_coverage",
            "field_definitions", "taxonomy_note", "monthly_derivation",
        )}
        for key, value in metadata.items():
            db.execute("INSERT INTO state_budget_metadata VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value_json=excluded.value_json", (key, _j(value)))
        for period, record in data["periods"].items():
            year, month = map(int, period.split("-"))
            end = f"{period}-{calendar.monthrange(year, month)[1]:02d}"
            db.execute("""INSERT INTO state_budget_periods VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
              ON CONFLICT(period) DO UPDATE SET period_start=excluded.period_start, period_end=excluded.period_end,
              period_semantics=excluded.period_semantics, classification=excluded.classification, source_unit=excluded.source_unit,
              source_format=excluded.source_format, source_file=excluded.source_file, reference_url=excluded.reference_url,
              archive_page_url=excluded.archive_page_url, field_order_json=excluded.field_order_json, source_record_json=excluded.source_record_json""",
              (period, f"{period}-01", end, record["period_semantics"], record["classification"], record["source_unit"],
               record["source_format"], record["source_file"], record["reference_url"], record.get("archive_page_url"),
               _j(record.get("field_order_in_source", [])), _j({k: v for k, v in record.items() if k != "categories"})))
            for category in record["categories"]:
                a = category["amounts"]
                db.execute("""INSERT INTO state_budget (period,level,code,source_code,parent_code,name_es,name_full_es,description_es,name,name_full,description,
                  initial_credits_thousand_eur,credit_modifications_thousand_eur,final_credits_thousand_eur,classification,source_url,source_file,source_format,source_row_json)
                  VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(period,level,code) DO UPDATE SET
                  source_code=excluded.source_code,parent_code=excluded.parent_code,name_es=excluded.name_es,name_full_es=excluded.name_full_es,description_es=excluded.description_es,
                  name=excluded.name,name_full=excluded.name_full,description=excluded.description,initial_credits_thousand_eur=excluded.initial_credits_thousand_eur,
                  credit_modifications_thousand_eur=excluded.credit_modifications_thousand_eur,final_credits_thousand_eur=excluded.final_credits_thousand_eur,
                  classification=excluded.classification,source_url=excluded.source_url,source_file=excluded.source_file,source_format=excluded.source_format,source_row_json=excluded.source_row_json""",
                  (period, category["level"], category["code"], category["source_code"], category.get("parent_code"), category["name_es"], category.get("name_full_es"), category.get("description_es"), category.get("name"), category.get("name_full"), category.get("description"), a.get(FIELDS[0]), a.get(FIELDS[1]), a.get(FIELDS[2]), record["classification"], record["reference_url"], record["source_file"], record["source_format"], _j(category.get("source_row"))))
                budget_id = db.execute("SELECT id FROM state_budget WHERE period=? AND level=? AND code=?", (period, category["level"], category["code"])).fetchone()[0]
                m = category.get("monthly_flows", {}); f = m.get("flows", {}); b = m.get("balance_changes", {})
                db.execute("""INSERT INTO state_expenditure (budget_id,period,level,code,committed_expenses_thousand_eur,recognized_obligations_thousand_eur,payments_thousand_eur,pending_payments_thousand_eur,credit_remnants_thousand_eur,
                  monthly_committed_expenses_thousand_eur,monthly_recognized_obligations_thousand_eur,monthly_payments_thousand_eur,monthly_credit_modifications_thousand_eur,monthly_pending_payments_change_thousand_eur,monthly_credit_remnants_change_thousand_eur,monthly_status,monthly_coverage_json,source_url)
                  VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(period,level,code) DO UPDATE SET budget_id=excluded.budget_id,committed_expenses_thousand_eur=excluded.committed_expenses_thousand_eur,recognized_obligations_thousand_eur=excluded.recognized_obligations_thousand_eur,payments_thousand_eur=excluded.payments_thousand_eur,pending_payments_thousand_eur=excluded.pending_payments_thousand_eur,credit_remnants_thousand_eur=excluded.credit_remnants_thousand_eur,monthly_committed_expenses_thousand_eur=excluded.monthly_committed_expenses_thousand_eur,monthly_recognized_obligations_thousand_eur=excluded.monthly_recognized_obligations_thousand_eur,monthly_payments_thousand_eur=excluded.monthly_payments_thousand_eur,monthly_credit_modifications_thousand_eur=excluded.monthly_credit_modifications_thousand_eur,monthly_pending_payments_change_thousand_eur=excluded.monthly_pending_payments_change_thousand_eur,monthly_credit_remnants_change_thousand_eur=excluded.monthly_credit_remnants_change_thousand_eur,monthly_status=excluded.monthly_status,monthly_coverage_json=excluded.monthly_coverage_json,source_url=excluded.source_url""",
                  (budget_id, period, category["level"], category["code"], a.get(FIELDS[3]), a.get(FIELDS[4]), a.get(FIELDS[5]), a.get(FIELDS[6]), a.get(FIELDS[7]), f.get(FLOW_FIELDS[1]), f.get(FLOW_FIELDS[2]), f.get(FLOW_FIELDS[3]), f.get(FLOW_FIELDS[0]), b.get("pending_payments_thousand_eur"), b.get("credit_remnants_thousand_eur"), m.get("status", "unavailable"), _j(m.get("category_coverage", {})), record["reference_url"]))
        if audit:
            db.execute("INSERT INTO state_budget_audit VALUES(1,?,datetime('now')) ON CONFLICT(id) DO UPDATE SET summary_json=excluded.summary_json, imported_at=excluded.imported_at", (_j(audit),))
        db.commit()


def update_budget_execution(
    project_root: Path,
    history_db: Path,
    *,
    data_path: Path | None = None,
    audit_path: Path | None = None,
) -> dict[str, int]:
    data_path = data_path or project_root / "data" / "budget_execution_history.json"
    audit_path = audit_path or project_root / "data" / "budget_execution_history_audit.summary.json"
    data = json.loads(data_path.read_text(encoding="utf-8"))
    current_year = datetime.now(UTC).year
    with httpx.Client(follow_redirects=True, timeout=90, headers={"User-Agent": "Paisometro budget updater/1.0"}) as client:
        sources = _discover_current_sources(client, current_year)
        latest_metadata = {}
        for period in sorted(data.get("periods", {}), reverse=True):
            record = data["periods"][period]
            for category in record.get("categories", []):
                latest_metadata.setdefault((category["level"], category["code"]), category)
        updated = 0
        for period, source in sorted(sources.items()):
            try:
                response = client.get(source["url"])
                response.raise_for_status()
                record = _parse_source(response.content, source["format"], period, source, latest_metadata)
                record.update({
                    "reference_url": source["url"], "archive_page_url": source["archive_url"],
                    "source_file": Path(unquote(urlparse(source["url"]).path)).name,
                    "source_format": source["format"],
                })
                record["validation"] = _validate(record)
                record["data_completeness"] = "complete" if record["validation"]["all_financial_columns_reconcile"] else "partial"
                data.setdefault("periods", {})[period] = record
                updated += 1
                LOGGER.info("Updated IGAE budget history period %s from %s", period, source["url"])
            except Exception as exc:
                LOGGER.warning("Keeping existing IGAE period %s; update failed: %s", period, exc)
    _derive_monthly(data["periods"])
    data["last_published_period_seen_on_current_index"] = max(data["periods"], default=None)
    data.setdefault("generation", {})["updated_at_utc"] = datetime.now(UTC).isoformat(timespec="seconds")
    data["generation"]["live_updated_periods"] = updated
    audit = json.loads(audit_path.read_text(encoding="utf-8")) if audit_path.exists() else {}
    _upsert_db(history_db, data, audit)
    content = json.dumps(data, ensure_ascii=False, indent=2) + "\n"
    fd, temporary = tempfile.mkstemp(prefix=f"{data_path.name}.", suffix=".tmp", dir=data_path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as handle:
            handle.write(content)
        os.replace(temporary, data_path)
    except Exception:
        if os.path.exists(temporary):
            os.unlink(temporary)
        raise
    return {"available_periods": len(data.get("periods", {})), "updated_periods": updated, "budget_rows": sum(len(r.get("categories", [])) for r in data.get("periods", {}).values())}
